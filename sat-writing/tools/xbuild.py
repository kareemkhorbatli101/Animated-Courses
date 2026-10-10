#!/usr/bin/env python3
"""Build the Book 3 .docx from data/exercises, then measure the rendered PDF.

The measuring is the point of the second half. Ten of the hundred and twenty
book-level checks are claims about the printed page -- that no exercise is split
across a page break, that ninety Arabic blocks were rendered, that the page count
falls inside its band -- and a claim about a page can only be tested on a page.
So this writes build/doc-stats.json from the PDF and the document's own XML, and
tools/xchecks.py reads it.
"""
import collections
import glob
import json
import os
import re
import subprocess
import sys
import zipfile

import yaml
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_TAB_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
SPEC = yaml.safe_load(open(os.path.join(ROOT, 'data', 'spec.yaml')))
CH = SPEC['chapters']
DOM = SPEC['domain_order']
DOMS = SPEC['domains']
RL = SPEC['rule_labels']
AR = SPEC['arabic']
LABELS = SPEC['exercise_rules']['labels']
SERIF = 'Georgia'
ARSERIF = 'FreeSerif'
NAME = 'Writing-the-Five-Fields'
# The two markers the measurement counts. They are ordinary lines of the book --
# a reader uses both -- and each appears exactly once where it belongs.
OPENER_MARK = 'Element %d of fifteen'
PART_MARK = 'Part %d.%d'
AR_NOTE_MARK = 'ملاحظة المجال'


def load():
    out = []
    for p in sorted(glob.glob(os.path.join(ROOT, 'data', 'exercises', '*.yaml'))):
        out.append(yaml.safe_load(open(p)))
    out.sort(key=lambda d: d['chapter'])
    return out


# ---------------------------------------------------------------------------
# page furniture
# ---------------------------------------------------------------------------
def style(doc):
    n = doc.styles['Normal']
    n.font.name = SERIF
    n.font.size = Pt(10.5)
    n._element.rPr.rFonts.set(qn('w:eastAsia'), SERIF)
    n.paragraph_format.space_after = Pt(0)
    n.paragraph_format.line_spacing = 1.08
    for s in doc.sections:
        s.page_width, s.page_height = Inches(8.5), Inches(11)
        s.left_margin = s.right_margin = Inches(0.95)
        s.top_margin = s.bottom_margin = Inches(0.85)


def para(doc, text='', size=10.5, bold=False, italic=False, align=None, before=0,
         after=0, indent=0, caps=False, grey=False, first=0, keep=False):
    p = doc.add_paragraph()
    r = p.add_run(text)
    r.font.size, r.bold, r.italic, r.font.name = Pt(size), bold, italic, SERIF
    if caps:
        r.font.all_caps = True
        r.font.spacing = Pt(1)
    if grey:
        r.font.color.rgb = RGBColor(0x66, 0x66, 0x66)
    if align:
        p.alignment = align
    p.paragraph_format.space_before = Pt(before)
    p.paragraph_format.space_after = Pt(after)
    if indent:
        p.paragraph_format.left_indent = Inches(indent)
    if first:
        p.paragraph_format.first_line_indent = Inches(first)
    if keep:
        p.paragraph_format.keep_together = True
    return p


def rule(doc, after=6, before=2, color='BBBBBB'):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(before)
    p.paragraph_format.space_after = Pt(after)
    pr = p._p.get_or_add_pPr()
    bd = OxmlElement('w:pBdr')
    b = OxmlElement('w:bottom')
    b.set(qn('w:val'), 'single')
    b.set(qn('w:sz'), '4')
    b.set(qn('w:space'), '1')
    b.set(qn('w:color'), color)
    bd.append(b)
    pr.append(bd)
    return p


def page_break(doc):
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)


def new_section(doc, header_text):
    s = doc.add_section(WD_SECTION.NEW_PAGE)
    s.page_width, s.page_height = Inches(8.5), Inches(11)
    s.left_margin = s.right_margin = Inches(0.95)
    s.top_margin = s.bottom_margin = Inches(0.85)
    s.header.is_linked_to_previous = False
    s.footer.is_linked_to_previous = False
    h = s.header.paragraphs[0]
    h.text = ''
    r = h.add_run(header_text)
    r.font.size = Pt(8.5)
    r.font.name = SERIF
    r.font.all_caps = True
    r.font.spacing = Pt(1)
    r.font.color.rgb = RGBColor(0x66, 0x66, 0x66)
    h.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    f = s.footer.paragraphs[0]
    f.text = ''
    fld = OxmlElement('w:fldSimple')
    fld.set(qn('w:instr'), 'PAGE')
    f._p.append(fld)
    f.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for rr in f.runs:
        rr.font.size = Pt(9)
        rr.font.name = SERIF
    return s


def arabic_para(doc, text, size=10.5, bold=False, before=0, after=3, indent=0.0):
    """A right-to-left paragraph in an Arabic-capable serif face."""
    p = doc.add_paragraph()
    pr = p._p.get_or_add_pPr()
    pr.append(OxmlElement('w:bidi'))
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    p.paragraph_format.space_before = Pt(before)
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.right_indent = Inches(indent)
    r = p.add_run(text)
    r.bold = bold
    r.font.size = Pt(size)
    r.font.name = ARSERIF
    rpr = r._element.get_or_add_rPr()
    rpr.append(OxmlElement('w:rtl'))
    fonts = rpr.find(qn('w:rFonts'))
    if fonts is None:
        fonts = OxmlElement('w:rFonts')
        rpr.insert(0, fonts)
    for a in ('w:cs', 'w:ascii', 'w:hAnsi', 'w:eastAsia'):
        fonts.set(qn(a), ARSERIF)
    sz = OxmlElement('w:szCs')
    sz.set(qn('w:val'), str(int(size * 2)))
    rpr.append(sz)
    return p


# ---------------------------------------------------------------------------
# one exercise
# ---------------------------------------------------------------------------
def emph(p, text, size=10.5, grey=False, caps=False, bold=False):
    """Add text to a paragraph, rendering *word* as italic.

    The rule index in data/spec.yaml marks the words a rule is about with
    asterisks -- "Two subjects joined by *and* take a plural verb" -- because a
    rule about the word *and* has to point at the word. On the page those have to
    be italics: the first build printed the asterisks.
    """
    for i, piece in enumerate(re.split(r'\*([^*]+)\*', text)):
        if not piece:
            continue
        r = p.add_run(piece)
        r.italic = bool(i % 2)
        r.bold = bold
        r.font.size = Pt(size)
        r.font.name = SERIF
        if caps:
            r.font.all_caps = True
            r.font.spacing = Pt(0.6)
        if grey:
            r.font.color.rgb = RGBColor(0x99, 0x99, 0x99)
    return p


def tag_line(doc, x):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(7)
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.tab_stops.add_tab_stop(Inches(6.6), WD_TAB_ALIGNMENT.RIGHT)
    r = p.add_run('%d' % x['n'])
    r.bold = True
    r.font.size = Pt(10)
    r.font.name = SERIF
    r2 = p.add_run('  %s' % x['id'])
    r2.font.size = Pt(7.5)
    r2.font.name = SERIF
    r2.font.color.rgb = RGBColor(0x99, 0x99, 0x99)
    p.add_run('\t')
    emph(p, '%s · %s · %s' % (RL[x['rule']][0], x['difficulty'],
                              SPEC['level_names'][x['level']]),
         size=7.5, grey=True, caps=True)
    return p


def table_block(doc, t):
    """The table of a quantitative exercise, as a real table."""
    cap = para(doc, t['title'], size=9, italic=True, indent=0.22, after=2, grey=True)
    tb = doc.add_table(rows=len(t['rows']) + 1, cols=len(t['cols']))
    tb.style = 'Table Grid'
    tb.autofit = True
    tb.left_indent = Inches(0.22)
    for j, c in enumerate(t['cols']):
        cell = tb.cell(0, j)
        cell.text = ''
        rr = cell.paragraphs[0].add_run(c)
        rr.bold = True
        rr.font.size = Pt(8.5)
        rr.font.name = SERIF
    for i, row in enumerate(t['rows'], 1):
        for j, c in enumerate(row):
            cell = tb.cell(i, j)
            cell.text = ''
            rr = cell.paragraphs[0].add_run(c)
            rr.font.size = Pt(8.5)
            rr.font.name = SERIF
    # A table is not a paragraph, so keep_with_next on the caption cannot reach past
    # it: the first build left ten exercises with their head line at the foot of one
    # page and the table overleaf. Marking every row unsplittable and every cell
    # paragraph keep-with-next pins the table to the text on both sides of it.
    for row in tb.rows:
        trpr = row._tr.get_or_add_trPr()
        trpr.append(OxmlElement('w:cantSplit'))
        for cell in row.cells:
            for pp in cell.paragraphs:
                pp.paragraph_format.keep_together = True
                pp.paragraph_format.keep_with_next = True
    return cap, tb


def one_exercise(doc, x):
    """Every paragraph of an exercise is pinned to the next, so the stem never
    stands at the foot of a page with its options overleaf. An exercise is at most
    a third of a page, so nothing has to break."""
    kept = [tag_line(doc, x)]
    ch = x['chapter']
    if ch <= 13:
        kept.append(para(doc, x['carrier'], size=10, indent=0.22, after=4))
    elif ch == 14:
        kept.append(para(doc, 'While researching a topic, a student has taken the '
                              'following notes.', size=9, indent=0.22, after=2,
                         grey=True))
        for note in x['notes']:
            b = doc.add_paragraph()
            b.paragraph_format.left_indent = Inches(0.42)
            b.paragraph_format.space_after = Pt(1)
            rr = b.add_run('•  ' + note)
            rr.font.size = Pt(9.5)
            rr.font.name = SERIF
            kept.append(b)
        gap = doc.add_paragraph()
        gap.paragraph_format.space_after = Pt(2)
        kept.append(gap)
    else:
        cap, tb = table_block(doc, x['table'])
        kept.append(cap)
        kept.append(para(doc, x['claim'], size=10, indent=0.22, before=3, after=4))
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Inches(0.22)
    p.paragraph_format.space_after = Pt(3)
    rs = p.add_run(x['stem'])
    rs.font.size = Pt(10)
    rs.font.name = SERIF
    kept.append(p)
    for i, o in enumerate(x['options']):
        op = doc.add_paragraph()
        op.paragraph_format.left_indent = Inches(0.52)
        op.paragraph_format.first_line_indent = Inches(-0.22)
        op.paragraph_format.space_after = Pt(1)
        rl = op.add_run('%s)  ' % LABELS[i])
        rl.font.size = Pt(10)
        rl.font.name = SERIF
        rl.font.color.rgb = RGBColor(0x55, 0x55, 0x55)
        ro = op.add_run(o)
        ro.font.size = Pt(10)
        ro.font.name = SERIF
        kept.append(op)
    for pp in kept:
        pp.paragraph_format.keep_together = True
    for pp in kept[:-1]:
        pp.paragraph_format.keep_with_next = True


# ---------------------------------------------------------------------------
# the body
# ---------------------------------------------------------------------------
def chapter_opener(doc, d):
    c = CH[d['chapter']]
    para(doc, OPENER_MARK % d['chapter'], size=9, caps=True, grey=True, after=4)
    para(doc, c['title'], size=22, bold=True, after=3)
    para(doc, 'Part %s · %s' % (c['part'], SPEC['book_parts'][c['part']]['title']),
         size=10, italic=True, grey=True, after=12)
    para(doc, c['gloss'], size=12, italic=True, after=12)
    para(doc, 'The rules this chapter drills', size=9, caps=True, grey=True, after=4)
    for r in c['rules']:
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Inches(0.3)
        p.paragraph_format.first_line_indent = Inches(-0.3)
        p.paragraph_format.space_after = Pt(3)
        p.paragraph_format.keep_together = True
        rr = p.add_run('—  ')
        rr.font.size = Pt(10)
        rr.font.name = SERIF
        emph(p, RL[r][0], size=10)
        if r in c['hardest']:
            r3 = p.add_run('   hardest')
            r3.font.size = Pt(7.5)
            r3.font.name = SERIF
            r3.font.all_caps = True
            r3.font.color.rgb = RGBColor(0x99, 0x99, 0x99)
    para(doc, 'How the wrong answers are built', size=9, caps=True, grey=True,
         before=12, after=4)
    para(doc, 'Every wrong answer in this chapter is one of these: '
              + ', '.join(m.replace('_', ' ') for m in c['moves'])
              + '. Each is a way of being wrong that the real test uses, and each is '
                'named for the exercise it belongs to in the answer key.',
         size=10, after=12)
    para(doc, 'The five parts that follow', size=9, caps=True, grey=True, after=4)
    para(doc, 'The same element in the five knowledge domains, ten exercises each, '
              'each part rising from Foundation to Stretch. The home domain of this '
              'element is %s, and its part carries every one of the hardest rules '
              'above.' % DOMS[c['home']], size=10, after=0)
    page_break(doc)
    arabic_page(doc, d)


def arabic_page(doc, d):
    c = CH[d['chapter']]
    arabic_para(doc, '%s' % AR['element_names'][c['element']], size=16, bold=True,
                after=8)
    for part in AR['chapter_parts']:
        txt = d['summary_ar'][part]
        arabic_para(doc, AR['chapter_headings'][part], size=11, bold=True, before=6,
                    after=2)
        arabic_para(doc, txt, size=10.5, after=4, indent=0.12)


def part_head(doc, d, part, idx):
    page_break(doc)
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.tab_stops.add_tab_stop(Inches(6.6), WD_TAB_ALIGNMENT.RIGHT)
    r = p.add_run(PART_MARK % (d['chapter'], idx))
    r.bold = True
    r.font.size = Pt(11)
    r.font.name = SERIF
    r2 = p.add_run('  %s' % DOMS[part['domain']])
    r2.font.size = Pt(11)
    r2.font.name = SERIF
    home = ' · home domain' if CH[d['chapter']]['home'] == part['domain'] else ''
    r3 = p.add_run('\tten exercises · Foundation to Stretch%s' % home)
    r3.font.size = Pt(8)
    r3.font.name = SERIF
    r3.font.color.rgb = RGBColor(0x88, 0x88, 0x88)
    para(doc, SPEC['domain_grammar'][part['domain']], size=9, italic=True, grey=True,
         after=3)
    kept = [arabic_para(doc, '%s: %s' % (AR_NOTE_MARK, part['note_ar']), size=9.5,
                        after=3)]
    kept.append(rule(doc, before=2, after=3, color='999999'))
    for pp in kept[:-1]:
        pp.paragraph_format.keep_together = True
        pp.paragraph_format.keep_with_next = True


def body(doc, data):
    for d in data:
        c = CH[d['chapter']]
        new_section(doc, 'Chapter %d · %s' % (d['chapter'], c['title']))
        chapter_opener(doc, d)
        for idx, part in enumerate(d['parts'], 1):
            part_head(doc, d, part, idx)
            for x in part['exercises']:
                one_exercise(doc, x)


# ---------------------------------------------------------------------------
# front matter, appendices, answer key
# ---------------------------------------------------------------------------
def front_matter(doc, data):
    for _ in range(5):
        doc.add_paragraph()
    para(doc, SPEC['title'].upper(), size=30, bold=True,
         align=WD_ALIGN_PARAGRAPH.CENTER, after=10)
    para(doc, SPEC['subtitle'], size=12, italic=True,
         align=WD_ALIGN_PARAGRAPH.CENTER, after=40)
    para(doc, 'Fifteen elements · five knowledge domains · ten exercises in every cell',
         size=11, align=WD_ALIGN_PARAGRAPH.CENTER, grey=True, after=6)
    para(doc, 'The third book of the series, after the vocabulary reader and the '
              'two hundred passages', size=11, italic=True,
         align=WD_ALIGN_PARAGRAPH.CENTER, grey=True)
    page_break(doc)

    para(doc, 'What this book is for', size=17, bold=True, after=8)
    for t in [
        'The Writing and Language half of the SAT asks about twenty-seven questions, and '
        'they fall into four groups: the conventions of Standard English, the boundaries '
        'between sentences, the expression of ideas, and the use of quantitative '
        'evidence. A student who has met the same rule twenty, thirty, fifty times stops '
        'reasoning about it and starts seeing it, which is the only state in which these '
        'questions can be answered at the pace the test sets.',
        'This book gives every one of fifteen elements fifty exercises: one part of ten '
        'in each of five knowledge domains, rising from Foundation to Stretch inside '
        'every part. The domains are not decoration. A subject-verb agreement question '
        'about a Latin plural in biology, about a collective noun in social science and '
        'about an apparatus list in physics are three different questions, and a student '
        'who has drilled only one of them has drilled a third of the skill.',
        'Nothing here is a reading passage. The reader before this one does that work. '
        'Here the sentence is as short as the rule allows at Foundation and as long as '
        'the test ever makes it at Stretch, and every question is about the words in the '
        'blank and nothing else.',
    ]:
        para(doc, t, size=11, after=8)

    para(doc, 'How the book is built', size=17, bold=True, before=10, after=8)
    para(doc, 'Fifteen chapters, one for each element. Five parts in every chapter, one '
              'for each knowledge domain. Ten exercises in every part. Seven hundred and '
              'fifty exercises in all, and every one of them carries its level, its '
              'difficulty, the rule it drills and \u2014 in the answer key \u2014 the name of '
              'the fault in each of its three wrong answers.', size=11, after=8)
    for pno in sorted(SPEC['book_parts']):
        bp = SPEC['book_parts'][pno]
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Inches(0.3)
        p.paragraph_format.first_line_indent = Inches(-0.3)
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.keep_together = True
        r = p.add_run('Part %d.  %s' % (pno, bp['title']))
        r.bold = True
        r.font.size = Pt(10.5)
        r.font.name = SERIF
        r2 = p.add_run('  \u2014  %s %s, in the %s domain of the test.'
                       % ('chapters' if len(bp['chapters']) > 1 else 'chapter',
                          ', '.join(str(i) for i in bp['chapters']), bp['domain']))
        r2.font.size = Pt(10.5)
        r2.font.name = SERIF

    para(doc, 'The four levels', size=17, bold=True, before=12, after=8)
    para(doc, 'Position inside a part fixes the level, so a student always knows how '
              'hard an exercise is meant to be before attempting it.', size=11, after=6)
    for lv in (1, 2, 3, 4):
        pos = [p for p in sorted(SPEC['level_by_pos']) if SPEC['level_by_pos'][p] == lv]
        lo, hi = SPEC['carrier_words'][lv]
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Inches(0.3)
        p.paragraph_format.first_line_indent = Inches(-0.3)
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.keep_together = True
        r = p.add_run('%s.' % SPEC['level_names'][lv])
        r.bold = True
        r.font.size = Pt(10.5)
        r.font.name = SERIF
        r2 = p.add_run('  Exercises %s of every part. The sentence runs %d to %d words '
                       'with the answer in place, and the words between the governing '
                       'word and the blank run %d to %d.'
                       % (' and '.join(str(i) for i in pos), lo, hi,
                          SPEC['span_gap'][lv][0], SPEC['span_gap'][lv][1]))
        r2.font.size = Pt(10.5)
        r2.font.name = SERIF

    para(doc, 'How to use it', size=17, bold=True, before=12, after=8)
    for t in [
        'Work a whole part of ten at a sitting. The ten rise, so the first two should '
        'feel easy and the last two should not, and stopping at the point where they '
        'stop being easy is the one way to learn nothing.',
        'Mark the answer before looking at the key, and then read the trap line even '
        'when you were right. The trap line names the wrong answer most students '
        'choose and says why it attracts, which is the part of an answer key worth '
        'reading twice.',
        'Every chapter opens with the rule in Arabic: the rule itself, how the test '
        'asks about it, the trap, and why it matters. Each part opens with a note in '
        'Arabic on what that knowledge domain does to this element. They are there so '
        'that a student who thinks about grammar in Arabic does not have to translate '
        'the explanation before using it.',
    ]:
        para(doc, t, size=11, after=8)


def contents(doc, data):
    new_section(doc, 'Contents')
    para(doc, 'Contents', size=20, bold=True, after=10)
    for pno in sorted(SPEC['book_parts']):
        bp = SPEC['book_parts'][pno]
        para(doc, 'Part %d · %s' % (pno, bp['title']), size=11, bold=True, before=10,
             after=4, keep=True)
        for i in bp['chapters']:
            c = CH[i]
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Inches(0.45)
            p.paragraph_format.first_line_indent = Inches(-0.45)
            p.paragraph_format.space_after = Pt(3)
            p.paragraph_format.keep_together = True
            r = p.add_run('%d.  ' % i)
            r.bold = True
            r.font.size = Pt(10.5)
            r.font.name = SERIF
            r2 = p.add_run(c['title'])
            r2.font.size = Pt(10.5)
            r2.font.name = SERIF
            r3 = p.add_run('  ·  fifty exercises  ·  home domain %s' % DOMS[c['home']])
            r3.font.size = Pt(8.5)
            r3.font.name = SERIF
            r3.font.color.rgb = RGBColor(0x99, 0x99, 0x99)
    para(doc, 'Appendix A · the grid of the book', size=10.5, before=12, after=3)
    para(doc, 'Appendix B · the five knowledge domains', size=10.5, after=3)
    para(doc, 'Appendix C · the rule index', size=10.5, after=3)
    para(doc, 'Appendix D · how the wrong answers are built', size=10.5, after=3)
    para(doc, 'Answer key · seven hundred and fifty rows', size=10.5, after=3)


def appendix_a(doc, data):
    new_section(doc, 'Appendix A · the grid of the book')
    para(doc, 'Appendix A', size=20, bold=True, after=2)
    para(doc, 'The grid of the book', size=12, italic=True, after=8)
    para(doc, 'Fifteen elements down the side, five knowledge domains across the top, '
              'ten exercises in every cell. The number in a cell is the exercise number '
              'the cell begins at.', size=10, grey=True, after=8)
    tb = doc.add_table(rows=16, cols=6)
    tb.style = 'Table Grid'
    hdr = ['Element'] + DOM
    for j, h in enumerate(hdr):
        cell = tb.cell(0, j)
        cell.text = ''
        r = cell.paragraphs[0].add_run(h)
        r.bold = True
        r.font.size = Pt(8)
        r.font.name = SERIF
    for i in range(1, 16):
        c = CH[i]
        cell = tb.cell(i, 0)
        cell.text = ''
        r = cell.paragraphs[0].add_run('%d. %s' % (i, c['title']))
        r.font.size = Pt(8)
        r.font.name = SERIF
        for j, d in enumerate(DOM, 1):
            cell = tb.cell(i, j)
            cell.text = ''
            r = cell.paragraphs[0].add_run(str((i - 1) * 50 + (j - 1) * 10 + 1))
            r.font.size = Pt(8)
            r.font.name = SERIF
            r.font.color.rgb = RGBColor(0x66, 0x66, 0x66)
    para(doc, 'Totals', size=9, caps=True, grey=True, before=12, after=4)
    tot = collections.Counter()
    for d in data:
        for part in d['parts']:
            for x in part['exercises']:
                tot[x['difficulty']] += 1
                tot['L%d' % x['level']] += 1
    para(doc, 'By difficulty: %d easy, %d medium, %d hard. By level: %d Foundation, '
              '%d Developing, %d Target, %d Stretch.'
              % (tot['easy'], tot['medium'], tot['hard'], tot['L1'], tot['L2'],
                 tot['L3'], tot['L4']), size=10, after=4)
    para(doc, 'Every part of ten holds the same shape: two Foundation, three '
              'Developing, three Target, two Stretch, and three easy, four medium, '
              'three hard. The shape is fixed by position, so the ladder is the same '
              'in all seventy-five parts.', size=10, after=0)


def appendix_b(doc, data):
    new_section(doc, 'Appendix B · the five knowledge domains')
    para(doc, 'Appendix B', size=20, bold=True, after=2)
    para(doc, 'The five knowledge domains', size=12, italic=True, after=8)
    para(doc, 'The same element is a different question in a different field, because '
              'each field has its own grammar. This is what each domain brings, and it '
              'is why the book is a grid and not a list.', size=10, grey=True, after=10)
    for d in DOM:
        para(doc, DOMS[d], size=12, bold=True, before=8, after=3, keep=True)
        para(doc, SPEC['domain_grammar'][d], size=10.5, indent=0.2, after=3)
        homes = [i for i in sorted(CH) if CH[i]['home'] == d]
        para(doc, 'Home domain of: %s.'
                  % ', '.join('%d %s' % (i, CH[i]['title'].lower()) for i in homes),
             size=9.5, indent=0.2, grey=True, after=3)


def appendix_c(doc, data):
    new_section(doc, 'Appendix C · the rule index')
    para(doc, 'Appendix C', size=20, bold=True, after=2)
    para(doc, 'The rule index', size=12, italic=True, after=8)
    para(doc, 'Every rule the book drills, with the chapter it belongs to and the '
              'number of exercises that rest on it. A student who gets a question wrong '
              'can find the rule here and then drill the other exercises that use it.',
         size=10, grey=True, after=10)
    used = collections.Counter()
    for d in data:
        for part in d['parts']:
            for x in part['exercises']:
                used[x['rule']] += 1
    for i in sorted(CH):
        c = CH[i]
        para(doc, '%d · %s' % (i, c['title']), size=10.5, bold=True, before=8, after=3,
             keep=True)
        for r in c['rules']:
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Inches(0.3)
            p.paragraph_format.first_line_indent = Inches(-0.3)
            p.paragraph_format.space_after = Pt(2)
            p.paragraph_format.keep_together = True
            emph(p, RL[r][0], size=9.5)
            r2 = p.add_run('   %s · %d exercises' % (r.replace('_', ' '), used[r]))
            r2.font.size = Pt(8)
            r2.font.name = SERIF
            r2.font.color.rgb = RGBColor(0x99, 0x99, 0x99)


def appendix_d(doc, data):
    new_section(doc, 'Appendix D · how the wrong answers are built')
    para(doc, 'Appendix D', size=20, bold=True, after=2)
    para(doc, 'How the wrong answers are built', size=12, italic=True, after=8)
    para(doc, 'Every wrong answer in the book carries the name of its fault, and the '
              'answer key gives that name. These are the faults, with the number of '
              'wrong answers built on each. Learning to recognize the fault is faster '
              'than learning to recognize the rule, and it transfers to questions this '
              'book does not contain.', size=10, grey=True, after=10)
    mv = collections.Counter()
    where = collections.defaultdict(set)
    for d in data:
        for part in d['parts']:
            for x in part['exercises']:
                for L, v in (x.get('faults') or {}).items():
                    mv[v['move']] += 1
                    where[v['move']].add(x['chapter'])
    for m, n in sorted(mv.items(), key=lambda kv: (-kv[1], kv[0])):
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Inches(0.3)
        p.paragraph_format.first_line_indent = Inches(-0.3)
        p.paragraph_format.space_after = Pt(3)
        p.paragraph_format.keep_together = True
        r = p.add_run(m.replace('_', ' '))
        r.bold = True
        r.font.size = Pt(10)
        r.font.name = SERIF
        r2 = p.add_run('   %d wrong answers · chapters %s'
                       % (n, ', '.join(str(i) for i in sorted(where[m]))))
        r2.font.size = Pt(8.5)
        r2.font.name = SERIF
        r2.font.color.rgb = RGBColor(0x99, 0x99, 0x99)


def answer_key(doc, data):
    new_section(doc, 'Answer key')
    para(doc, 'Answer key', size=20, bold=True, after=2)
    para(doc, 'Seven hundred and fifty rows, in the order of the book', size=12,
         italic=True, after=6)
    para(doc, 'Each row gives the key, the rule it rests on, what makes it right, and '
              'the wrong answer most likely to attract, with the name of its fault. The '
              'trap line is the one worth reading twice.', size=10, grey=True, after=10)
    for d in data:
        c = CH[d['chapter']]
        for part in d['parts']:
            h = doc.add_paragraph()
            h.paragraph_format.space_before = Pt(9)
            h.paragraph_format.space_after = Pt(3)
            h.paragraph_format.keep_with_next = True
            rh = h.add_run('%d · %s' % (d['chapter'], c['title']))
            rh.bold = True
            rh.font.size = Pt(10.5)
            rh.font.name = SERIF
            rh2 = h.add_run('   %s' % DOMS[part['domain']])
            rh2.font.size = Pt(8.5)
            rh2.font.name = SERIF
            rh2.font.color.rgb = RGBColor(0x99, 0x99, 0x99)
            for x in part['exercises']:
                p = doc.add_paragraph()
                p.paragraph_format.left_indent = Inches(0.62)
                p.paragraph_format.first_line_indent = Inches(-0.62)
                p.paragraph_format.space_after = Pt(2)
                p.paragraph_format.keep_together = True
                r1 = p.add_run('%s  %s  ' % (x['id'], x['key']))
                r1.bold = True
                r1.font.size = Pt(9)
                r1.font.name = SERIF
                emph(p, RL[x['rule']][0] + '   ', size=8, grey=True)
                r3 = p.add_run(x['why'])
                r3.font.size = Pt(9)
                r3.font.name = SERIF
                r4 = p.add_run('  Trap: ' + x['trap'])
                r4.font.size = Pt(9)
                r4.font.name = SERIF
                r4.italic = True
                r4.font.color.rgb = RGBColor(0x55, 0x55, 0x55)
                r5 = p.add_run('  [%s]' % '; '.join(
                    '%s %s' % (L, v['move'].replace('_', ' '))
                    for L, v in sorted((x.get('faults') or {}).items())))
                r5.font.size = Pt(7.5)
                r5.font.name = SERIF
                r5.font.color.rgb = RGBColor(0x99, 0x99, 0x99)


# ---------------------------------------------------------------------------
# measurement
# ---------------------------------------------------------------------------
ID = re.compile(r'C\d\d-[A-Z]{3}-E\d\d')


def measure(docx, pdf, data):
    S = {}
    xml = ''
    with zipfile.ZipFile(docx) as z:
        xml = ''.join(z.read(n).decode('utf8', 'replace') for n in z.namelist()
                      if n.startswith('word/') and n.endswith('.xml'))
    # Arabic is counted in the document's own XML rather than the PDF. A PDF text
    # layer of right-to-left script does not survive matching: pdftotext finds the
    # one-word heading and loses the three-word ones, so counting phrases there
    # measures the converter rather than the book. The XML text is exact.
    paras = [re.sub(r'<[^>]+>', '', m) for m in re.findall(r'<w:p[ >].*?</w:p>', xml,
                                                           re.S)]
    names = {AR['element_names'][CH[i]['element']] for i in CH}
    heads = set(AR['chapter_headings'].values())
    S['arabic_blocks'] = (sum(1 for t in paras if t.strip() in names)
                          + sum(1 for t in paras if t.strip().startswith(AR_NOTE_MARK)))
    S['arabic_headings'] = sum(1 for t in paras if t.strip() in heads)
    txt = ''
    if os.path.exists(pdf):
        txt = subprocess.run(['pdftotext', '-layout', pdf, '-'],
                             capture_output=True, text=True, timeout=600).stdout
    BIDI = {ord(c): None for c in '‎‏‪‫‬‭‮'
                                  '⁦⁧⁨⁩'}
    txt = txt.translate(BIDI)
    pages = [p for p in txt.split('\f')]
    if pages and not pages[-1].strip():
        pages.pop()
    S['pages'] = len(pages)
    # The opener marker prints in small capitals, so the match ignores case.
    S['openers'] = sum(1 for i in sorted(CH)
                       if re.search(re.escape(OPENER_MARK % i), txt, re.I))
    S['part_heads'] = sum(1 for i in sorted(CH) for j in range(1, 6)
                          if re.search(re.escape(PART_MARK % (i, j)) + r'\b', txt))
    keypos = txt.find('Seven hundred and fifty rows')
    body_txt = txt[:keypos] if keypos > 0 else txt
    key_txt = txt[keypos:] if keypos > 0 else ''
    S['exercises'] = len(set(ID.findall(body_txt)))
    S['key_rows'] = len(set(ID.findall(key_txt)))
    S['appendices'] = len(set(re.findall(r'Appendix ([A-D])\b', txt)))
    # The contents runs from its own heading to the first chapter opener, however
    # many pages that takes.
    a = txt.find('Contents')
    b = txt.find(OPENER_MARK % 1, a if a > 0 else 0)
    cont = txt[a:b] if 0 < a < b else ''
    if not cont:
        b2 = txt.upper().find((OPENER_MARK % 1).upper(), a if a > 0 else 0)
        cont = txt[a:b2] if 0 < a < b2 else ''
    S['contents'] = sum(1 for i in sorted(CH) if CH[i]['title'] in cont)
    # An exercise is split when its own stretch of a page does not hold all four of
    # its option labels: the head of one exercise to the head of the next is one
    # exercise's worth of text, and it has to be on one page.
    split = 0
    for pg in pages[:keypos if keypos < 0 else None]:
        if keypos > 0 and 'Seven hundred and fifty rows' in pg:
            break
        marks = list(ID.finditer(pg))
        for k, m in enumerate(marks):
            chunk = pg[m.start():marks[k + 1].start() if k + 1 < len(marks) else len(pg)]
            if not all('%s)' % L in chunk for L in LABELS):
                split += 1
    S['split'] = split
    S['arabic_pages'] = sum(1 for pg in pages
                            if any('؀' <= ch <= 'ۿ' for ch in pg))
    return S


def main():
    data = load()
    doc = Document()
    style(doc)
    front_matter(doc, data)
    contents(doc, data)
    body(doc, data)
    appendix_a(doc, data)
    appendix_b(doc, data)
    appendix_c(doc, data)
    appendix_d(doc, data)
    answer_key(doc, data)
    out = os.path.join(ROOT, 'build', '%s.docx' % NAME)
    os.makedirs(os.path.dirname(out), exist_ok=True)
    doc.save(out)
    print('wrote %s  %.0f KB' % (out, os.path.getsize(out) / 1024))
    if '--no-pdf' not in sys.argv:
        subprocess.run(['libreoffice', '--headless', '--convert-to', 'pdf',
                        '--outdir', os.path.dirname(out), out],
                       capture_output=True, timeout=1800)
    pdf = os.path.join(ROOT, 'build', '%s.pdf' % NAME)
    S = measure(out, pdf, data)
    with open(os.path.join(ROOT, 'build', 'doc-stats.json'), 'w') as fh:
        json.dump(S, fh, indent=1, sort_keys=True)
    for k in sorted(S):
        print('  %-14s %s' % (k, S[k]))


if __name__ == '__main__':
    main()
