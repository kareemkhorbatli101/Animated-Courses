#!/usr/bin/env python3
"""Build the reader .docx from the passage files."""
import os, re, sys, glob, yaml
from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT, WD_BREAK
from docx.enum.section import WD_SECTION
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
SPEC = yaml.safe_load(open(os.path.join(ROOT, 'data', 'spec.yaml')))
STRANDS = yaml.safe_load(open(os.path.join(ROOT, 'data', 'strands.yaml')))
SERIF = 'Georgia'


def load():
    out = []
    for p in sorted(glob.glob(os.path.join(ROOT, 'data', 'passages', '*.yaml'))):
        d = yaml.safe_load(open(p))
        for x in d['passages']:
            x['field'], x['level'] = d['field'], d['level']
            x['paras'] = [' '.join(q.split()) for q in x['passage'].strip().split('\n\n') if q.strip()]
            out.append(x)
    order = {f: i for i, f in enumerate(SPEC['field_order'])}
    out.sort(key=lambda x: (x['level'], order[x['field']], x['strand']))
    for i, x in enumerate(out, 1):
        x['n'] = i
    return out


def style(doc):
    n = doc.styles['Normal']
    n.font.name = SERIF
    n.font.size = Pt(11)
    n._element.rPr.rFonts.set(qn('w:eastAsia'), SERIF)
    n.paragraph_format.space_after = Pt(0)
    n.paragraph_format.line_spacing = 1.1
    for s in doc.sections:
        s.page_width, s.page_height = Inches(8.5), Inches(11)
        s.left_margin = s.right_margin = Inches(1.0)
        s.top_margin = s.bottom_margin = Inches(0.9)


def para(doc, text='', size=11, bold=False, italic=False, align=None, before=0, after=0,
         indent=0, caps=False, grey=False, first=0):
    p = doc.add_paragraph()
    r = p.add_run(text)
    r.font.size, r.bold, r.italic, r.font.name = Pt(size), bold, italic, SERIF
    if caps:
        r.font.all_caps = True
        r.font.spacing = Pt(1)
    if grey:
        r.font.color.rgb = RGBColor(0x66, 0x66, 0x66)
    if align:
        p.alignment = align
    p.paragraph_format.space_before = Pt(before)
    p.paragraph_format.space_after = Pt(after)
    if indent:
        p.paragraph_format.left_indent = Inches(indent)
    if first:
        p.paragraph_format.first_line_indent = Inches(first)
    return p


def rule(doc, after=6, before=2, color='BBBBBB'):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(before)
    p.paragraph_format.space_after = Pt(after)
    pr = p._p.get_or_add_pPr()
    bd = OxmlElement('w:pBdr')
    b = OxmlElement('w:bottom')
    b.set(qn('w:val'), 'single'); b.set(qn('w:sz'), '4')
    b.set(qn('w:space'), '1'); b.set(qn('w:color'), color)
    bd.append(b); pr.append(bd)
    return p


def page_break(doc):
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)


def new_section(doc, header_text):
    s = doc.add_section(WD_SECTION.NEW_PAGE)
    s.page_width, s.page_height = Inches(8.5), Inches(11)
    s.left_margin = s.right_margin = Inches(1.0)
    s.top_margin = s.bottom_margin = Inches(0.9)
    s.header.is_linked_to_previous = False
    s.footer.is_linked_to_previous = False
    h = s.header.paragraphs[0]; h.text = ''
    r = h.add_run(header_text)
    r.font.size = Pt(8.5); r.font.name = SERIF
    r.font.all_caps = True; r.font.spacing = Pt(1)
    r.font.color.rgb = RGBColor(0x66, 0x66, 0x66)
    h.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    f = s.footer.paragraphs[0]; f.text = ''
    fld = OxmlElement('w:fldSimple'); fld.set(qn('w:instr'), 'PAGE')
    f._p.append(fld)
    f.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for rr in f.runs:
        rr.font.size = Pt(9); rr.font.name = SERIF
    return s


def anchor_box(doc, x):
    rule(doc, before=10, after=4, color='999999')
    para(doc, 'What this passage is for', size=9, caps=True, grey=True, after=3)
    for a in x['anchors']:
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Inches(0.22)
        p.paragraph_format.space_after = Pt(2)
        p.paragraph_format.keep_together = True
        r = p.add_run('—  ' + a['take'])
        r.font.size = Pt(10); r.font.name = SERIF
    rule(doc, before=4, after=0, color='999999')


def passage_page(doc, x, first_in_section=False):
    if not first_in_section:
        page_break(doc)
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(1)
    p.paragraph_format.tab_stops.add_tab_stop(Inches(6.5), WD_TAB_ALIGNMENT.RIGHT)
    r = p.add_run('%d' % x['n'])
    r.bold = True; r.font.size = Pt(10); r.font.name = SERIF
    r2 = p.add_run('\t%s · %s · %s' % (x['field'], x['strand'], x['move']))
    r2.font.size = Pt(8); r2.font.name = SERIF
    r2.font.color.rgb = RGBColor(0x88, 0x88, 0x88)
    para(doc, x['title'], size=15, bold=True, after=8)
    for i, q in enumerate(x['paras']):
        para(doc, q, size=11, after=7, first=0 if i == 0 else 0.2)
    anchor_box(doc, x)


ARSERIF = 'FreeSerif'
QSETS = {}
AR = SPEC['arabic']
SLOTS = SPEC['slots']


def load_questions():
    out = {}
    for p in sorted(glob.glob(os.path.join(ROOT, 'data', 'questions', '*.yaml'))):
        d = yaml.safe_load(open(p))
        for s in d['sets']:
            out[s['id']] = s
    return out


def disp(s):
    return str(s).replace('«', '“').replace('»', '”')


def arabic_para(doc, text, size=10.5, bold=False, before=0, after=3, indent=0.0):
    """A right-to-left paragraph in an Arabic-capable serif face."""
    p = doc.add_paragraph()
    pr = p._p.get_or_add_pPr()
    bidi = OxmlElement('w:bidi')
    pr.append(bidi)
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    p.paragraph_format.space_before = Pt(before)
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.right_indent = Inches(indent)
    r = p.add_run(text)
    r.bold = bold
    r.font.size = Pt(size)
    r.font.name = ARSERIF
    rpr = r._element.get_or_add_rPr()
    rtl = OxmlElement('w:rtl')
    rpr.append(rtl)
    fonts = rpr.find(qn('w:rFonts'))
    if fonts is None:
        fonts = OxmlElement('w:rFonts')
        rpr.insert(0, fonts)
    for a in ('w:cs', 'w:ascii', 'w:hAnsi', 'w:eastAsia'):
        fonts.set(qn(a), ARSERIF)
    sz = OxmlElement('w:szCs')
    sz.set(qn('w:val'), str(int(size * 2)))
    rpr.append(sz)
    return p


def question_label(q):
    return '%s · %s · %s' % (q['type'].replace('_', ' '), q['difficulty'],
                                       SLOTS[q['slot']]['domain'])


def one_question(doc, n, q):
    # Every paragraph of a question is collected and then pinned together, so that a
    # stem never appears at the foot of a page with its four options overleaf. A
    # question is at most about a third of a page, so nothing is forced to break.
    kept = []
    head = doc.add_paragraph()
    head.paragraph_format.space_before = Pt(7)
    head.paragraph_format.space_after = Pt(2)
    head.paragraph_format.keep_with_next = True
    kept.append(head)
    head.paragraph_format.tab_stops.add_tab_stop(Inches(6.5), WD_TAB_ALIGNMENT.RIGHT)
    r = head.add_run('%d.%d' % (n, q['slot']))
    r.bold = True
    r.font.size = Pt(10)
    r.font.name = SERIF
    r2 = head.add_run('\t' + question_label(q))
    r2.font.size = Pt(7.5)
    r2.font.name = SERIF
    r2.font.all_caps = True
    r2.font.spacing = Pt(0.6)
    r2.font.color.rgb = RGBColor(0x99, 0x99, 0x99)

    if q['type'] == 'cross_text':
        kept.append(para(doc, disp(q['sibling_gloss']), size=9.5, indent=0.22, after=4,
                         italic=True, grey=True))
    if q['type'] == 'synthesis':
        kept.append(para(doc, 'While researching a topic, a student has taken the following '
                              'notes.', size=9.5, indent=0.22, after=2, grey=True))
        for note in q['notes']:
            b = doc.add_paragraph()
            b.paragraph_format.left_indent = Inches(0.42)
            b.paragraph_format.space_after = Pt(1)
            rr = b.add_run('•  ' + disp(note))
            rr.font.size = Pt(9.5)
            rr.font.name = SERIF
            kept.append(b)
        gap = doc.add_paragraph()
        gap.paragraph_format.space_after = Pt(2)
        kept.append(gap)
    if q.get('carrier'):
        kept.append(para(doc, disp(q['carrier']), size=10, indent=0.22, after=4))

    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Inches(0.22)
    p.paragraph_format.space_after = Pt(3)
    rs = p.add_run(disp(q['stem']))
    rs.font.size = Pt(10.5)
    rs.font.name = SERIF
    kept.append(p)

    for i, o in enumerate(q['options']):
        op = doc.add_paragraph()
        op.paragraph_format.left_indent = Inches(0.52)
        op.paragraph_format.first_line_indent = Inches(-0.22)
        op.paragraph_format.space_after = Pt(1)
        op.paragraph_format.keep_together = True
        rl = op.add_run('%s)  ' % SPEC['question_rules']['labels'][i])
        rl.font.size = Pt(10)
        rl.font.name = SERIF
        rl.font.color.rgb = RGBColor(0x55, 0x55, 0x55)
        ro = op.add_run(disp(o))
        ro.font.size = Pt(10)
        ro.font.name = SERIF
        kept.append(op)
    for pp in kept:
        pp.paragraph_format.keep_together = True
    for pp in kept[:-1]:
        pp.paragraph_format.keep_with_next = True


def question_pages(doc, x, qset):
    page_break(doc)
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.tab_stops.add_tab_stop(Inches(6.5), WD_TAB_ALIGNMENT.RIGHT)
    r = p.add_run('Questions %d.1 to %d.10' % (x['n'], x['n']))
    r.bold = True
    r.font.size = Pt(11)
    r.font.name = SERIF
    r2 = p.add_run('\ton passage %d, %s' % (x['n'], x['title']))
    r2.font.size = Pt(8.5)
    r2.font.name = SERIF
    r2.font.color.rgb = RGBColor(0x88, 0x88, 0x88)
    rule(doc, before=2, after=2, color='999999')
    for q in qset['questions']:
        one_question(doc, x['n'], q)
    arabic_summary(doc, x, qset)


def arabic_summary(doc, x, qset):
    a = qset.get('summary_ar') or {}
    if not a:
        return
    kept = [rule(doc, before=11, after=3, color='999999'),
            arabic_para(doc, 'ملخّص المقطع %d بالعربية' % x['n'], size=10, bold=True, after=4)]
    for part in AR['parts']:
        if not a.get(part):
            continue
        kept.append(arabic_para(doc, '%s: %s' % (AR['part_headings'][part], a[part]),
                                size=9.5, after=3, indent=0.12))
    kept.append(rule(doc, before=3, after=0, color='999999'))
    for pp in kept[:-1]:
        pp.paragraph_format.keep_together = True
        pp.paragraph_format.keep_with_next = True


def answer_key(doc, items, qsets):
    new_section(doc, 'Appendix E · answer key')
    para(doc, 'Appendix E', size=20, bold=True, after=2)
    para(doc, 'Answer key', size=12, italic=True, after=4)
    para(doc, 'Two thousand questions in passage order. Each row gives the key, what makes it '
              'right, and the wrong answer most likely to attract. The trap line is the one '
              'worth reading twice.', size=10, grey=True, after=10)
    for x in items:
        qs = qsets.get(x['id'])
        if not qs:
            continue
        h = doc.add_paragraph()
        h.paragraph_format.space_before = Pt(10)
        h.paragraph_format.space_after = Pt(3)
        h.paragraph_format.keep_with_next = True
        rh = h.add_run('%d · %s' % (x['n'], x['title']))
        rh.bold = True
        rh.font.size = Pt(10.5)
        rh.font.name = SERIF
        rh2 = h.add_run('   %s · %s · %s' % (x['field'], x['strand'], x['move']))
        rh2.font.size = Pt(8)
        rh2.font.name = SERIF
        rh2.font.color.rgb = RGBColor(0x99, 0x99, 0x99)
        for q in qs['questions']:
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Inches(0.55)
            p.paragraph_format.first_line_indent = Inches(-0.55)
            p.paragraph_format.space_after = Pt(2)
            p.paragraph_format.keep_together = True
            r1 = p.add_run('%d.%-2d  %s  ' % (x['n'], q['slot'], q['key']))
            r1.bold = True
            r1.font.size = Pt(9.5)
            r1.font.name = SERIF
            r2 = p.add_run('%s · %s   ' % (q['type'].replace('_', ' '), q['difficulty']))
            r2.font.size = Pt(8)
            r2.font.name = SERIF
            r2.font.color.rgb = RGBColor(0x99, 0x99, 0x99)
            r3 = p.add_run(disp(q['why']))
            r3.font.size = Pt(9.5)
            r3.font.name = SERIF
            r4 = p.add_run('  Trap: ' + disp(q['trap']))
            r4.font.size = Pt(9.5)
            r4.font.name = SERIF
            r4.italic = True
            r4.font.color.rgb = RGBColor(0x55, 0x55, 0x55)


def front_matter(doc, items):
    for _ in range(5):
        doc.add_paragraph()
    para(doc, SPEC['title'].upper(), size=30, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, after=10)
    para(doc, SPEC['subtitle'], size=13, italic=True, align=WD_ALIGN_PARAGRAPH.CENTER, after=40)
    para(doc, 'Two hundred passages · five fields · fifty strands · four levels',
         size=11, align=WD_ALIGN_PARAGRAPH.CENTER, grey=True, after=6)
    para(doc, 'A companion to %s' % SPEC['companion'], size=11, italic=True,
         align=WD_ALIGN_PARAGRAPH.CENTER, grey=True)
    page_break(doc)

    para(doc, 'What this book is for', size=17, bold=True, after=8)
    for t in [
        'The SAT does not test knowledge of history or science. It does test reading, and it '
        'does so with short passages drawn from history, social studies, the humanities, '
        'literature and science. A student who has met the subject before reads such a '
        'passage faster, guesses unknown words better, and is not stopped by a name.',
        'That is not a claim about intelligence. It is the most reliable finding in the '
        'research on reading comprehension. Recht and Leslie, working with sixty-four '
        'junior-high students in 1988, found that weak readers who knew baseball recalled '
        'more of a passage about a baseball game than strong readers who did not. The study '
        'has been over-read and criticized, and the direction of the effect has held up: a '
        'systematic review of twenty-three studies in 2021 found more domain knowledge '
        'produced better comprehension of text in that domain.',
        'For a reader working in a second language there is a condition attached. Background '
        'knowledge pays only once enough of the words are known, a point the research '
        'literature calls the linguistic threshold. Work on second-language listening finds '
        'that learners below a certain vocabulary level get no benefit at all from knowing '
        'the topic. The usable version of that finding is the coverage figure attributed to '
        'Hu and Nation: about ninety-eight per cent of the running words need to be known '
        'for unassisted reading, and about ninety-five per cent for minimally acceptable '
        'comprehension.',
        'So this book has to do two things at once. It has to raise knowledge, and it has to '
        'stay inside the threshold while doing it. Ninety-eight per cent of three hundred '
        'words is where the arithmetic lands: at most six words per passage may fall outside '
        "the level's assumed vocabulary, and at the first two levels every one of them "
        'is explained where it stands.',
    ]:
        para(doc, t, after=8)
    para(doc, 'The figures above reached this book through secondary sources. Clarke, Hu and '
              'Nation, and Recht and Leslie should be checked at source before any of them is '
              'quoted in a classroom.', size=10, italic=True, grey=True, before=6)
    page_break(doc)

    para(doc, 'How the book is built', size=17, bold=True, after=8)
    para(doc, 'Five fields, ten strands in each, four levels in each strand. A strand is one '
              'thread of content followed through all four levels, so each strand receives '
              'twelve hundred words of treatment and can be read downward as well as across.',
         after=10)
    para(doc, 'The four levels do four different jobs, in the same order every time. This is '
              'what makes the coverage graded rather than merely longer.', after=8)
    for lv in (1, 2, 3, 4):
        L = SPEC['levels'][lv]
        move = SPEC['moves'][lv]
        p = doc.add_paragraph(); p.paragraph_format.space_after = Pt(7)
        r = p.add_run('Level %d, %s — %s. ' % (lv, L['name'], move))
        r.bold = True; r.font.size = Pt(11); r.font.name = SERIF
        r2 = p.add_run(SPEC['move_notes'][move] + ' ' +
                       'Sentences average %.0f to %.0f words; vocabulary is assumed to the top %s '
                       'most frequent word forms.' % (L['sent_mean_min'], L['sent_mean_max'],
                                                      format(L['known_band'], ',d')))
        r2.font.size = Pt(11); r2.font.name = SERIF
    para(doc, 'So one strand in biology runs: a reindeer herd that reached six thousand and then '
              'forty-two (phenomenon), the growth curves and the lag that let a population pass '
              "its own ceiling (mechanism), the Hudson's Bay fur returns and what a trapping ledger "
              'can and cannot show (evidence), and the experiment that found food and predators '
              'acting together rather than separately (dispute).', before=10, after=8)
    para(doc, 'Every one of the two hundred target words from %s appears somewhere in these '
              'passages, in the same field and at the same level, used in running prose. A '
              'student can meet a word here and meet it again as a question.' % SPEC['companion'],
         after=8)
    page_break(doc)

    para(doc, 'How to use it', size=17, bold=True, after=8)
    for t in [
        'Read across or read down. Reading across means taking all ten passages of one field '
        'at one level, which gives a survey of the field in three thousand words. Reading '
        'down means taking one strand through its four levels, which shows the same subject '
        'treated four ways and is the better use of the book for anyone with time.',
        'Each passage ends with three lines headed what this passage is for. Read them after '
        'the passage, not before. If none of the three is a surprise, the passage has done '
        'its work and can be left. If one of them is, the passage is worth a second reading, '
        'and the sentence it came from is worth finding.',
        'Do not look up every unfamiliar word. The passages are built so that no more than '
        'six words in three hundred fall outside what a reader at that level is assumed to '
        'know, and the technical ones are explained in the sentence where they appear. A '
        'reader who stops at every name will never find out that the names do not matter.',
        'Keep a list of the field terms rather than of the hard words. Appendix B collects '
        'every term the book explains, with its explanation and the passage it appears in, '
        'and that list is the actual content of the book.',
    ]:
        para(doc, t, after=8)
    para(doc, 'What this book is not', size=13, bold=True, before=8, after=6)
    para(doc, 'These passages are not past papers and they are not a syllabus. The College '
              'Board publishes no list of subjects and no reading list, so the five fields and '
              'the fifty strands are a curricular judgment, informed by what the test is '
              'described as covering. The passages were written from general knowledge and '
              'checked for standard figures; a reviewer who intends to assert any particular '
              'fact in a classroom should verify it first.', after=8)
    page_break(doc)

    para(doc, 'Contents', size=17, bold=True, after=10)
    n = 1
    for lv in (1, 2, 3, 4):
        L = SPEC['levels'][lv]
        para(doc, 'Level %d — %s · %s' % (lv, L['name'], SPEC['moves'][lv]),
             size=12, bold=True, before=8, after=4)
        for f in SPEC['field_order']:
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Inches(0.25)
            p.paragraph_format.space_after = Pt(2)
            p.paragraph_format.tab_stops.add_tab_stop(Inches(6.5), WD_TAB_ALIGNMENT.RIGHT)
            r = p.add_run('%s\tpassages %d–%d' % (SPEC['fields'][f], n, n + 9))
            r.font.size = Pt(10.5); r.font.name = SERIF
            n += 10
    for t in ['The fifty strands', 'Appendix A — strand index',
              'Appendix B — the field terms explained',
              'Appendix C — where the vocabulary words appear',
              'Appendix D — reading log']:
        para(doc, t, size=12, bold=True, before=8, after=2)


def strand_map(doc):
    new_section(doc, 'The fifty strands')
    para(doc, 'The fifty strands', size=20, bold=True, after=4)
    para(doc, 'Each line is one strand. The four entries under it are the four levels, in order: '
              'phenomenon, mechanism, evidence, dispute.', size=10.5, italic=True, after=10)
    for f in SPEC['field_order']:
        para(doc, SPEC['fields'][f], size=13, bold=True, before=10, after=4)
        for code in sorted(STRANDS[f]):
            s = STRANDS[f][code]
            p = doc.add_paragraph()
            p.paragraph_format.space_after = Pt(1)
            p.paragraph_format.keep_together = True
            r = p.add_run('%s  %s' % (code, s['title']))
            r.bold = True; r.font.size = Pt(10.5); r.font.name = SERIF
            for k in ('l1', 'l2', 'l3', 'l4'):
                q = doc.add_paragraph()
                q.paragraph_format.left_indent = Inches(0.35)
                q.paragraph_format.space_after = Pt(0)
                q.paragraph_format.keep_together = True
                rr = q.add_run(' '.join(s[k].split()))
                rr.font.size = Pt(9.5); rr.font.name = SERIF
                rr.font.color.rgb = RGBColor(0x44, 0x44, 0x44)
            doc.add_paragraph().paragraph_format.space_after = Pt(3)


def body(doc, items):
    for lv in (1, 2, 3, 4):
        L = SPEC['levels'][lv]
        new_section(doc, 'Level %d · %s' % (lv, L['name']))
        for _ in range(3):
            doc.add_paragraph()
        para(doc, 'Level %d' % lv, size=11, caps=True, grey=True, after=2)
        para(doc, L['name'], size=28, bold=True, after=4)
        para(doc, SPEC['moves'][lv], size=14, italic=True, grey=True, after=12)
        rule(doc, after=12)
        para(doc, SPEC['move_notes'][SPEC['moves'][lv]], size=11.5, after=12)
        rows = [('Passage length', '%d to %d words' % (L['words_min'], L['words_max'])),
                ('Sentences', 'averaging %.0f to %.0f words' % (L['sent_mean_min'], L['sent_mean_max'])),
                ('Vocabulary assumed', 'the top %s most frequent word forms' % format(L['known_band'], ',d')),
                ('Words above that', 'at most %d, %s' % (L['above_band_max'],
                 'each explained in place' if L['gloss_all'] else 'explained or inferable')),
                ('Clue to the reader', re.sub(r'\s+', ' ', L['description']).strip())]
        for k, v in rows:
            p = doc.add_paragraph()
            p.paragraph_format.space_after = Pt(4)
            p.paragraph_format.left_indent = Inches(0.25)
            r = p.add_run('%s — ' % k); r.bold = True; r.font.size = Pt(10.5); r.font.name = SERIF
            r2 = p.add_run(v); r2.font.size = Pt(10.5); r2.font.name = SERIF
        for f in SPEC['field_order']:
            sub = [x for x in items if x['level'] == lv and x['field'] == f]
            new_section(doc, 'Level %d · %s · %s' % (lv, L['name'], SPEC['fields'][f]))
            para(doc, SPEC['fields'][f], size=18, bold=True, after=2)
            para(doc, 'Passages %d to %d' % (sub[0]['n'], sub[-1]['n']), size=10, grey=True, after=10)
            rule(doc, after=0)
            for i, x in enumerate(sub):
                passage_page(doc, x, first_in_section=(i == 0))
                if x['id'] in QSETS:
                    question_pages(doc, x, QSETS[x['id']])


def appendices(doc, items):
    by_id = {x['id']: x for x in items}
    new_section(doc, 'Appendix A · strand index')
    para(doc, 'Appendix A', size=20, bold=True, after=2)
    para(doc, 'Strand index', size=12, italic=True, after=4)
    para(doc, 'The four numbers are the passages for that strand, in level order.',
         size=10, grey=True, after=10)
    for f in SPEC['field_order']:
        para(doc, SPEC['fields'][f], size=12, bold=True, before=8, after=3)
        for code in sorted(STRANDS[f]):
            ns = []
            for lv in (1, 2, 3, 4):
                k = '%s-%s-L%d' % (f, code, lv)
                ns.append(str(by_id[k]['n']) if k in by_id else '-')
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Inches(0.25)
            p.paragraph_format.space_after = Pt(1)
            p.paragraph_format.tab_stops.add_tab_stop(Inches(6.5), WD_TAB_ALIGNMENT.RIGHT)
            r = p.add_run('%s  %s\t%s' % (code, STRANDS[f][code]['title'], '  ·  '.join(ns)))
            r.font.size = Pt(10); r.font.name = SERIF

    new_section(doc, 'Appendix B · the field terms explained')
    para(doc, 'Appendix B', size=20, bold=True, after=2)
    para(doc, 'The field terms explained', size=12, italic=True, after=4)
    para(doc, 'Every term this book defines, with the passage in which it is defined.',
         size=10, grey=True, after=10)
    terms = {}
    for x in items:
        for t in (x.get('terms') or []):
            terms.setdefault(t['term'].lower(), (t['term'], t['gloss'], x['n']))
    for k in sorted(terms):
        term, gloss, n = terms[k]
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(3)
        p.paragraph_format.keep_together = True
        r = p.add_run('%s  ' % term); r.bold = True; r.font.size = Pt(10.5); r.font.name = SERIF
        r2 = p.add_run(gloss); r2.font.size = Pt(10.5); r2.font.name = SERIF
        r3 = p.add_run('  (%d)' % n); r3.font.size = Pt(9); r3.font.name = SERIF
        r3.font.color.rgb = RGBColor(0x77, 0x77, 0x77)

    new_section(doc, 'Appendix C · where the vocabulary words appear')
    para(doc, 'Appendix C', size=20, bold=True, after=2)
    para(doc, 'Where the vocabulary words appear', size=12, italic=True, after=4)
    para(doc, 'The two hundred target words of %s, each with the passage in this book that uses '
              'it in context.' % SPEC['companion'], size=10, grey=True, after=10)
    rows = []
    for x in items:
        for v in (x.get('vocab_link') or []):
            rows.append((v, x['n']))
    rows.sort(key=lambda r: r[0].lower())
    per = (len(rows) + 2) // 3
    t = doc.add_table(rows=1, cols=3)
    for ci in range(3):
        cell = t.rows[0].cells[ci]
        cell.paragraphs[0].text = ''
        for w, n in rows[ci * per:(ci + 1) * per]:
            p = cell.add_paragraph()
            p.paragraph_format.space_after = Pt(0)
            r = p.add_run('%s  %d' % (w, n))
            r.font.size = Pt(9.5); r.font.name = SERIF
        if cell.paragraphs and not cell.paragraphs[0].runs and len(cell.paragraphs) > 1:
            cell.paragraphs[0]._p.getparent().remove(cell.paragraphs[0]._p)

    new_section(doc, 'Appendix D · reading log')
    para(doc, 'Appendix D', size=20, bold=True, after=2)
    para(doc, 'Reading log', size=12, italic=True, after=4)
    para(doc, 'Twenty rows of ten. Mark a passage when the three lines at the end of it hold no '
              'surprises.', size=10, grey=True, after=10)
    for lv in (1, 2, 3, 4):
        para(doc, 'Level %d — %s' % (lv, SPEC['levels'][lv]['name']),
             size=11, bold=True, before=8, after=3)
        tb = doc.add_table(rows=6, cols=11)
        tb.style = 'Table Grid'
        tb.autofit = False
        tb.allow_autofit = False
        tblPr = tb._tbl.tblPr
        for tag in ('w:tblLayout', 'w:tblW'):
            for el in tblPr.findall(qn(tag)):
                tblPr.remove(el)
        lay = OxmlElement('w:tblLayout'); lay.set(qn('w:type'), 'fixed'); tblPr.append(lay)
        tw = OxmlElement('w:tblW'); tw.set(qn('w:type'), 'dxa')
        tw.set(qn('w:w'), str(int(6.5 * 1440))); tblPr.append(tw)
        widths = [Inches(1.5)] + [Inches(0.5)] * 10
        grid = tb._tbl.find(qn('w:tblGrid'))
        for gc, w in zip(grid.findall(qn('w:gridCol')), widths):
            gc.set(qn('w:w'), str(int(w.twips)))
        for row in tb.rows:
            row.height = Inches(0.26)
            for cell, w in zip(row.cells, widths):
                cell.width = w
        for ri, f in enumerate(SPEC['field_order'], start=1):
            c = tb.cell(ri, 0); c.text = ''
            r = c.paragraphs[0].add_run(SPEC['fields'][f])
            r.font.size = Pt(8); r.bold = True; r.font.name = SERIF
        base = (lv - 1) * 50
        for ci in range(1, 11):
            c = tb.cell(0, ci); c.text = ''
            r = c.paragraphs[0].add_run('S%02d' % ci)
            r.font.size = Pt(8); r.bold = True; r.font.name = SERIF
        for ri in range(1, 6):
            for ci in range(1, 11):
                c = tb.cell(ri, ci); c.text = ''
                r = c.paragraphs[0].add_run(str(base + (ri - 1) * 10 + ci))
                r.font.size = Pt(8); r.font.name = SERIF
                r.font.color.rgb = RGBColor(0x99, 0x99, 0x99)


def main():
    items = load()
    QSETS.update(load_questions())
    doc = Document()
    style(doc)
    front_matter(doc, items)
    strand_map(doc)
    body(doc, items)
    appendices(doc, items)
    if QSETS:
        answer_key(doc, items, QSETS)
    out = os.path.join(ROOT, 'build', 'Reading-the-Five-Fields.docx')
    doc.save(out)
    print('wrote', out, '%.0f KB' % (os.path.getsize(out) / 1024))


if __name__ == '__main__':
    main()
