#!/usr/bin/env python3
"""Build the .docx from the item files. Questions and key come from one source,
so they cannot drift apart."""
import os, re, glob, yaml
from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT, WD_BREAK
from docx.enum.section import WD_SECTION
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
SPEC = yaml.safe_load(open(os.path.join(ROOT, 'data', 'spec.yaml')))
BLANK = SPEC['blank_token']
LET = 'ABCD'
SERIF = 'Georgia'


def load():
    items = []
    for p in sorted(glob.glob(os.path.join(ROOT, 'data', 'items', '*.yaml'))):
        d = yaml.safe_load(open(p))
        for it in d['items']:
            it['level'], it['domain'] = d['level'], d['domain']
            it['passage'] = re.sub(r'\s+', ' ', it['passage']).strip()
            it['pathway'] = re.sub(r'\s+', ' ', it['pathway']).strip()
            it['clues'] = [it['clue']] if isinstance(it['clue'], str) else list(it['clue'])
            items.append(it)
    order = {d: i for i, d in enumerate(SPEC['domain_order'])}
    items.sort(key=lambda i: (i['level'], order[i['domain']], i['id']))
    for n, it in enumerate(items, 1):
        it['n'] = n
    return items


def style(doc):
    n = doc.styles['Normal']
    n.font.name = SERIF
    n.font.size = Pt(11)
    n._element.rPr.rFonts.set(qn('w:eastAsia'), SERIF)
    pf = n.paragraph_format
    pf.space_after = Pt(0)
    pf.line_spacing = 1.08
    for s in doc.sections:
        s.page_width, s.page_height = Inches(8.5), Inches(11)
        s.left_margin = s.right_margin = Inches(0.95)
        s.top_margin = Inches(0.9)
        s.bottom_margin = Inches(0.9)


def para(doc, text='', size=11, bold=False, italic=False, align=None,
         before=0, after=0, indent=0, first=0, caps=False, grey=False):
    p = doc.add_paragraph()
    r = p.add_run(text)
    r.font.size = Pt(size)
    r.bold = bold
    r.italic = italic
    r.font.name = SERIF
    if caps:
        r.font.all_caps = True
        r.font.spacing = Pt(1)
    if grey:
        r.font.color.rgb = RGBColor(0x66, 0x66, 0x66)
    if align:
        p.alignment = align
    p.paragraph_format.space_before = Pt(before)
    p.paragraph_format.space_after = Pt(after)
    if indent:
        p.paragraph_format.left_indent = Inches(indent)
    if first:
        p.paragraph_format.first_line_indent = Inches(first)
    return p


def rule(doc, after=6):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(after)
    pr = p._p.get_or_add_pPr()
    bd = OxmlElement('w:pBdr')
    b = OxmlElement('w:bottom')
    b.set(qn('w:val'), 'single'); b.set(qn('w:sz'), '4')
    b.set(qn('w:space'), '1'); b.set(qn('w:color'), 'BBBBBB')
    bd.append(b); pr.append(bd)


def page_break(doc):
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)


def new_section(doc, header_text):
    s = doc.add_section(WD_SECTION.NEW_PAGE)
    s.page_width, s.page_height = Inches(8.5), Inches(11)
    s.left_margin = s.right_margin = Inches(0.95)
    s.top_margin = Inches(0.9); s.bottom_margin = Inches(0.9)
    s.header.is_linked_to_previous = False
    s.footer.is_linked_to_previous = False
    h = s.header.paragraphs[0]
    h.text = ''
    r = h.add_run(header_text)
    r.font.size = Pt(8.5); r.font.name = SERIF
    r.font.all_caps = True; r.font.spacing = Pt(1)
    r.font.color.rgb = RGBColor(0x66, 0x66, 0x66)
    h.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    f = s.footer.paragraphs[0]
    f.text = ''
    fr = f.add_run()
    fld = OxmlElement('w:fldSimple'); fld.set(qn('w:instr'), 'PAGE')
    f._p.append(fld)
    f.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for rr in f.runs:
        rr.font.size = Pt(9); rr.font.name = SERIF
    return s


def item_block(doc, it, with_code=True):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(10)
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.keep_with_next = True
    p.paragraph_format.keep_together = True
    p.paragraph_format.tab_stops.add_tab_stop(Inches(6.6), WD_TAB_ALIGNMENT.RIGHT)
    r = p.add_run('%d.' % it['n']); r.bold = True; r.font.size = Pt(11); r.font.name = SERIF
    if with_code:
        r2 = p.add_run('\tL%d · %s · %02d' % (it['level'], it['domain'], (it['n'] - 1) % 10 + 1))
        r2.font.size = Pt(8); r2.font.name = SERIF
        r2.font.color.rgb = RGBColor(0x88, 0x88, 0x88)
    pp = para(doc, it['passage'], size=11, after=4, indent=0.28)
    pp.paragraph_format.keep_with_next = True
    pp.paragraph_format.keep_together = True
    q = doc.add_paragraph()
    q.paragraph_format.left_indent = Inches(0.28)
    q.paragraph_format.space_after = Pt(2)
    q.paragraph_format.keep_together = True
    for x in (1.90, 3.45, 5.00):
        q.paragraph_format.tab_stops.add_tab_stop(Inches(x), WD_TAB_ALIGNMENT.LEFT)
    for i, o in enumerate(it['options']):
        if i:
            q.add_run('\t').font.size = Pt(10.5)
        r = q.add_run('(%s) %s' % (LET[i], o))
        r.font.size = Pt(10.5); r.font.name = SERIF
    rule(doc, after=0)


def front_matter(doc, items):
    for _ in range(4):
        doc.add_paragraph()
    para(doc, SPEC['title'].upper(), size=30, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, after=10)
    para(doc, SPEC['subtitle'], size=13, italic=True, align=WD_ALIGN_PARAGRAPH.CENTER, after=40)
    para(doc, '200 questions · four levels · five subject areas',
         size=11, align=WD_ALIGN_PARAGRAPH.CENTER, grey=True, after=6)
    para(doc, 'with a full answer key explaining every wrong option',
         size=11, align=WD_ALIGN_PARAGRAPH.CENTER, grey=True)
    page_break(doc)

    para(doc, 'What this book is', size=16, bold=True, after=8)
    for t in [
        'Every question in this book tests one skill: working out what a word must mean '
        'from the sentences around it. There is no grammar here, no punctuation and no '
        'sentence structure. Two hundred times, you are given a short passage with one '
        'word removed and four words to choose from.',
        'This is the commonest question type in the Reading and Writing section of the '
        'digital SAT. The College Board calls it Words in Context, and it sits in the '
        'Craft and Structure domain, which is about twenty-eight per cent of that '
        'section. The questions here are built to the same shape: one short passage, one '
        'blank, four single-word options.',
        'The book is for a student who reads English well but did not grow up in it. That '
        'student usually loses marks on these questions for one reason: the passage '
        'contains a name, a date or a technical detail that is unfamiliar, and the eye '
        'stops there. The whole design of this book is aimed at that habit.',
    ]:
        para(doc, t, after=8)
    para(doc, 'How the levels work', size=13, bold=True, before=10, after=6)
    for lv in (1, 2, 3, 4):
        L = SPEC['levels'][lv]
        p = doc.add_paragraph(); p.paragraph_format.space_after = Pt(6)
        r = p.add_run('Level %d — %s. ' % (lv, L['name'])); r.bold = True
        r.font.name = SERIF; r.font.size = Pt(11)
        r2 = p.add_run('%s. Target words from the %s. Passages of %d to %d words. Clue %s.'
                       % (re.sub(r'\s+', ' ', L['description']).strip().rstrip('.'),
                          L['band'], L['words_min'], L['words_max'], L['clue_distance']))
        r2.font.name = SERIF; r2.font.size = Pt(11)
    page_break(doc)

    para(doc, 'How a context question works', size=16, bold=True, after=8)
    para(doc, 'Here is a question of the kind you will meet at Level 2.', after=8)
    ex = ('Although the treaty was signed in 1783, the border it described remained '
          '________ for decades: both governments went on sending surveyors into the '
          'same disputed valleys, and each produced maps the other would not accept.')
    para(doc, ex, indent=0.28, after=4, italic=True)
    para(doc, '(A) renowned     (B) contested     (C) adjacent     (D) temporary',
         indent=0.28, after=10)
    for h, t in [
        ('Step 1 — read for the clue, not for the facts.',
         'The words that decide the answer are "the same disputed valleys". Everything '
         'else in the passage — 1783, the treaty, the surveyors, the maps — is true, '
         'interesting, and irrelevant to the choice.'),
        ('Step 2 — say what the blank must mean before you look at the options.',
         'A border that two governments are still arguing over is a border that is '
         'argued over. Cover the four words and describe the meaning in your own '
         'language if that is easier. If you cannot describe it, you have not found the '
         'clue yet, and the options will not help you find it.'),
        ('Step 3 — test every option against the clue, not against your ear.',
         '(A) renowned means famous; nothing here concerns fame. (C) adjacent means next '
         'to, which is true of every border and therefore tells you nothing. (D) '
         'temporary is contradicted by "for decades". (B) contested restates "disputed". '
         'The answer is (B).'),
        ('Step 4 — check that the whole passage still makes sense with your word in place.',
         'This takes four seconds and catches most careless mistakes.'),
    ]:
        p = doc.add_paragraph(); p.paragraph_format.space_after = Pt(8)
        r = p.add_run(h + ' '); r.bold = True; r.font.name = SERIF; r.font.size = Pt(11)
        r2 = p.add_run(t); r2.font.name = SERIF; r2.font.size = Pt(11)
    page_break(doc)

    para(doc, 'The four traps', size=16, bold=True, after=8)
    para(doc, 'Every wrong option in this book is wrong for a stated reason, and the '
              'answer key names the reason by its tag. There are four.', after=10)
    for t, d in SPEC['traps'].items():
        p = doc.add_paragraph(); p.paragraph_format.space_after = Pt(8)
        r = p.add_run('%s  ' % t); r.bold = True; r.font.name = SERIF; r.font.size = Pt(11)
        r2 = p.add_run(d); r2.font.name = SERIF; r2.font.size = Pt(11)
    para(doc, 'Names, dates and numbers are almost never the clue',
         size=13, bold=True, before=12, after=6)
    for t in [
        'This is the most useful sentence in the book, so it is worth reading twice. In '
        'two hundred questions, the word that licenses the answer is a common English '
        'word — disputed, gone, only, still, never, instead. It is not a proper noun and '
        'it is not a figure.',
        'A passage may mention the Grand Banks, a bench mark, Sowerby, the Admiralty or '
        'a chronometer. You are not expected to know what any of them are. When your eye '
        'stops on an unfamiliar name, mark it in your head as scenery and read on to the '
        'end of the sentence. The clue will be in ordinary words, and it will usually be '
        'within one sentence of the blank.',
        'Each level puts more of this scenery in your way on purpose. By Level 4 about '
        'half of what you read in a passage has no bearing on the answer, and two '
        'separate clues have to be put together before the blank can be filled. That is '
        'the skill the test is measuring.',
    ]:
        para(doc, t, after=8)
    page_break(doc)

    para(doc, 'How to use the book', size=16, bold=True, after=8)
    for t in [
        'Work in sets of ten — one subject area at one level — and mark them before you '
        'go on. Ten questions take about twelve minutes at Level 1 and twenty at Level 4. '
        'On the real test you will have about seventy seconds a question, but speed is '
        'the last thing to train, not the first.',
        'Read the key for every question you got right as well as every one you got '
        'wrong. The key gives the clue, the reasoning in one sentence, and a note on each '
        'of the three wrong options. The three wrong options are where the learning is: '
        'if you can say why (C) is wrong, you will not be caught by (C) again.',
        'Keep the words you miss in a list of your own, with the sentence you met them '
        'in. A word learned inside a sentence is remembered; a word learned from a '
        'glossary is not.',
        'Appendix C has a box for each of the twenty sets. Fill it in. The pattern tells '
        'you more than the total: a student who scores eight at Level 3 in history and '
        'four at Level 3 in physical science has a vocabulary problem in one subject, '
        'not a reading problem in general.',
    ]:
        para(doc, t, after=8)
    para(doc, 'A note on what this book is not', size=13, bold=True, before=10, after=6)
    para(doc, 'These are practice questions written to the published shape of the digital '
              'SAT. They are not past papers, and no book outside the College Board can '
              'promise you the exact difficulty of a live test. Use the levels as a '
              'ladder, not as a score prediction.', after=8)
    page_break(doc)

    para(doc, 'Contents', size=16, bold=True, after=10)
    n = 1
    for lv in (1, 2, 3, 4):
        L = SPEC['levels'][lv]
        para(doc, 'Chapter %d — Level %d, %s' % (lv, lv, L['name']),
             size=12, bold=True, before=8, after=4)
        for dom in SPEC['domain_order']:
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Inches(0.28)
            p.paragraph_format.space_after = Pt(2)
            p.paragraph_format.tab_stops.add_tab_stop(Inches(6.6), WD_TAB_ALIGNMENT.RIGHT)
            r = p.add_run('%s\tquestions %d–%d' % (SPEC['domains'][dom], n, n + 9))
            r.font.size = Pt(10.5); r.font.name = SERIF
            n += 10
    for t in ['Answer key', 'Appendix A — the two hundred words, alphabetically',
              'Appendix B — the words by subject area', 'Appendix C — score tracker']:
        para(doc, t, size=12, bold=True, before=8, after=2)


def questions(doc, items):
    for lv in (1, 2, 3, 4):
        L = SPEC['levels'][lv]
        new_section(doc, 'Chapter %d · %s' % (lv, L['name']))
        para(doc, 'Chapter %d' % lv, size=11, caps=True, grey=True, after=2)
        para(doc, 'Level %d — %s' % (lv, L['name']), size=22, bold=True, after=10)
        rule(doc, after=10)
        para(doc, 'What changes at this level', size=12, bold=True, after=6)
        rows = [('Target words', L['band']),
                ('Passage length', '%d to %d words' % (L['words_min'], L['words_max'])),
                ('Where the clue is', L['clue_distance']),
                ('How it reads', re.sub(r'\s+', ' ', L['description']).strip())]
        for k, v in rows:
            p = doc.add_paragraph()
            p.paragraph_format.space_after = Pt(4)
            p.paragraph_format.left_indent = Inches(0.28)
            r = p.add_run('%s — ' % k); r.bold = True; r.font.size = Pt(10.5); r.font.name = SERIF
            r2 = p.add_run(v); r2.font.size = Pt(10.5); r2.font.name = SERIF
        para(doc, 'Fifty questions: ten in each of the five subject areas. Answers begin on '
                  'the first page of the answer key.', before=10, size=10.5, italic=True)
        for dom in SPEC['domain_order']:
            sub = [i for i in items if i['level'] == lv and i['domain'] == dom]
            new_section(doc, 'Chapter %d · %s · %s' % (lv, L['name'], SPEC['domains'][dom]))
            para(doc, SPEC['domains'][dom], size=15, bold=True, after=2)
            para(doc, 'Questions %d–%d' % (sub[0]['n'], sub[-1]['n']),
                 size=10, grey=True, after=4)
            rule(doc, after=2)
            for it in sub:
                item_block(doc, it)


def key(doc, items):
    new_section(doc, 'Answer key')
    para(doc, 'Answer key', size=22, bold=True, after=8)
    para(doc, 'Each entry gives the answer, the words in the passage that license it, the '
              'reasoning in one sentence, and a note on each wrong option with the trap it '
              'belongs to. T1 register match, logic fail. T2 near-synonym, wrong '
              'collocation or degree. T3 the familiar sense of a word whose rarer sense is '
              'being tested. T4 reversed direction.', size=10.5, italic=True, after=10)
    cur = None
    for it in items:
        tag = (it['level'], it['domain'])
        if tag != cur:
            cur = tag
            para(doc, 'Level %d — %s — %s' % (it['level'], SPEC['levels'][it['level']]['name'],
                                              SPEC['domains'][it['domain']]),
                 size=12, bold=True, before=14, after=6)
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(8)
        p.paragraph_format.space_after = Pt(2)
        p.paragraph_format.keep_with_next = True
        r = p.add_run('%d.  (%s) %s' % (it['n'], it['answer'].upper(), it['word']))
        r.bold = True; r.font.size = Pt(11); r.font.name = SERIF
        clue_txt = ' — and — '.join('“%s”' % re.sub(r'\s+', ' ', c) for c in it['clues'])
        p2 = doc.add_paragraph()
        p2.paragraph_format.left_indent = Inches(0.28)
        p2.paragraph_format.space_after = Pt(2)
        p2.paragraph_format.keep_with_next = True
        r = p2.add_run('The clue is %s.' % clue_txt)
        r.font.size = Pt(10.5); r.font.name = SERIF
        r2 = p2.add_run('  (%s)' % it['relation'])
        r2.font.size = Pt(9); r2.italic = True; r2.font.name = SERIF
        r2.font.color.rgb = RGBColor(0x77, 0x77, 0x77)
        pw = para(doc, it['pathway'], size=10.5, indent=0.28, after=3)
        pw.paragraph_format.keep_with_next = True
        wrong_left = sum(1 for i, o in enumerate(it['options']) if 'abcd'[i] != it['answer'])
        done = 0
        for i, o in enumerate(it['options']):
            L = 'abcd'[i]
            if L == it['answer']:
                continue
            p3 = doc.add_paragraph()
            p3.paragraph_format.left_indent = Inches(0.28)
            p3.paragraph_format.space_after = Pt(1)
            done += 1
            p3.paragraph_format.keep_with_next = (done < wrong_left)
            p3.paragraph_format.keep_together = True
            r = p3.add_run('(%s) %s — ' % (LET[i], o))
            r.font.size = Pt(10.5); r.font.name = SERIF; r.italic = True
            r2 = p3.add_run(it['key'][L])
            r2.font.size = Pt(10.5); r2.font.name = SERIF
            r3 = p3.add_run('  (%s)' % it['traps'][L])
            r3.font.size = Pt(9); r3.font.name = SERIF
            r3.font.color.rgb = RGBColor(0x77, 0x77, 0x77)


def appendices(doc, items):
    new_section(doc, 'Appendix A · the two hundred words')
    para(doc, 'Appendix A', size=20, bold=True, after=2)
    para(doc, 'The two hundred words, alphabetically', size=12, italic=True, after=8)
    para(doc, 'Words appear in the form the question uses. The figure after each word '
              'is the number of that question.', size=10, grey=True, after=8)
    rows = sorted(((it['word'], it['n'], it['level']) for it in items), key=lambda r: r[0].lower())
    t = doc.add_table(rows=0, cols=4)
    for i in range(0, len(rows), 1):
        pass
    per = (len(rows) + 3) // 4
    cols = [rows[i*per:(i+1)*per] for i in range(4)]
    tr = t.add_row()
    for ci, col in enumerate(cols):
        cell = tr.cells[ci]
        cell.paragraphs[0].text = ''
        for w, n, lv in col:
            p = cell.add_paragraph()
            p.paragraph_format.space_after = Pt(0)
            r = p.add_run('%s  %d' % (w, n))
            r.font.size = Pt(9); r.font.name = SERIF
    for row in t.rows:
        for c in row.cells:
            if c.paragraphs and not c.paragraphs[0].runs and len(c.paragraphs) > 1:
                c.paragraphs[0]._p.getparent().remove(c.paragraphs[0]._p)

    new_section(doc, 'Appendix B · the words by subject area')
    para(doc, 'Appendix B', size=20, bold=True, after=2)
    para(doc, 'The words by subject area', size=12, italic=True, after=10)
    for dom in SPEC['domain_order']:
        para(doc, SPEC['domains'][dom], size=12, bold=True, before=10, after=4)
        for lv in (1, 2, 3, 4):
            sub = [i for i in items if i['domain'] == dom and i['level'] == lv]
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Inches(0.28)
            p.paragraph_format.space_after = Pt(3)
            r = p.add_run('Level %d  ' % lv); r.bold = True
            r.font.size = Pt(10); r.font.name = SERIF
            r2 = p.add_run(' · '.join('%s (%d)' % (i['word'], i['n']) for i in sub))
            r2.font.size = Pt(10); r2.font.name = SERIF

    new_section(doc, 'Appendix C · score tracker')
    para(doc, 'Appendix C', size=20, bold=True, after=2)
    para(doc, 'Score tracker', size=12, italic=True, after=8)
    para(doc, 'One box for each set of ten. Write the score and the date. Read down a '
              'column to see a subject area across the levels; read across a row to see '
              'one level across the subjects.', size=10.5, after=10)
    t = doc.add_table(rows=5, cols=6)
    t.style = 'Table Grid'
    hdr = ['', ] + [SPEC['domains'][d] for d in SPEC['domain_order']]
    for ci, h in enumerate(hdr):
        c = t.cell(0, ci); c.text = ''
        r = c.paragraphs[0].add_run(h)
        r.bold = True; r.font.size = Pt(9); r.font.name = SERIF
    for ri, lv in enumerate((1, 2, 3, 4), start=1):
        c = t.cell(ri, 0); c.text = ''
        r = c.paragraphs[0].add_run('Level %d\n%s' % (lv, SPEC['levels'][lv]['name']))
        r.bold = True; r.font.size = Pt(9); r.font.name = SERIF
        for ci in range(1, 6):
            cell = t.cell(ri, ci); cell.text = ''
            p = cell.paragraphs[0]
            r = p.add_run('\n     /10\n\ndate:\n')
            r.font.size = Pt(9); r.font.name = SERIF
            r.font.color.rgb = RGBColor(0x88, 0x88, 0x88)
    para(doc, 'Two hundred questions in all. A set below six out of ten is worth doing '
              'again after two weeks, not immediately: the second attempt tests memory '
              'of the answer, the third tests the skill.', size=10.5, before=12)


def main():
    items = load()
    doc = Document()
    style(doc)
    front_matter(doc, items)
    questions(doc, items)
    key(doc, items)
    appendices(doc, items)
    out = os.path.join(ROOT, 'build', 'SAT-Words-in-Context.docx')
    doc.save(out)
    print('wrote', out, '%.0f KB' % (os.path.getsize(out) / 1024))


if __name__ == '__main__':
    main()
